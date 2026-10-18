from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from datetime import date

wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F); GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True); H1 = Font(name=F, bold=True, size=14); H2 = Font(name=F, bold=True, size=11, color="FFFFFF")
YEL = PatternFill("solid", fgColor="FFFF00"); NAVY = PatternFill("solid", fgColor="1B2430"); GREY = PatternFill("solid", fgColor="E9E4D8")
GBP = '£#,##0;(£#,##0);-'; PCT = '0.0%'; NUM = '#,##0;(#,##0);-'; NUM1 = '#,##0.0;(#,##0.0);-'
thin = Side(style="thin", color="BBBBBB")

def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, color=c.font.color, size=c.font.size)

# ---------------- Assumptions ----------------
A = wb.active; A.title = "Assumptions"
A.column_dimensions["A"].width = 46; A.column_dimensions["B"].width = 14; A.column_dimensions["C"].width = 14; A.column_dimensions["D"].width = 70
A["A1"] = "Jobs Tracker — 36-month financial model"; A["A1"].font = H1
A["A2"] = "Blue cells with yellow fill are inputs; everything else is a formula. Month 1 = November 2026 (Gaswise billing starts). Figures in £, excluding VAT."
A["A2"].font = Font(name=F, italic=True)

rows = [
    ("COMMERCIAL", None, None, None),
    ("Model start (month 1)", date(2026, 11, 1), None, "Gaswise billing and set-up fee start 1 November 2026"),
    ("Average new contract, £ per month at list", 1150, None, "Franck: target mix of 15–30, 30–50 and 50+ worker firms at £30/worker + £50/admin; CBES-sized firm is £1,100"),
    ("Set-up fee per new firm, £ one-off", 4000, None, "Rate card, October 2026; every new firm pays it"),
    ("Monthly logo churn", 0.01, None, "Assumption: ~11% annual; no data yet (product live 3 months)"),
    ("Price uplift at each annual renewal", 0.03, None, "Applied to the whole subscription base at months 13 and 25"),
    ("Existing MRR at month 0 (CBES net £440 + Gaswise £2,600), £", 3040, None, "CBES £1,100 list less 60% first-year pilot discount; Gaswise £2,600 at list"),
    ("CBES discount ends: month", 10, None, "Initial term 1 Aug 2026 – 31 Jul 2027; list price from August 2027 = month 10"),
    ("CBES MRR step-up when discount ends, £", 660, None, "£1,100 − £440"),
    ("Existing firms at month 0", 2, None, "CBES and Gaswise"),
    ("Sales achievement vs plan (sensitivity)", 1.0, None, "1 = the ramp as printed on the deck; 0.5 = half the new firms each month. Use it to see when cash runs out if sales come slower"),
    ("ADD-ONS (storage packs, SMS credits, AI quote-writing)", None, None, None),
    ("Add-ons as share of subscription MRR, year 1", 0.03, None, "Adoption ~30% of firms in year 1; AI at £5 per field worker, storage £20 per 50 GB"),
    ("Add-ons as share of subscription MRR, year 2", 0.10, None, ""),
    ("Add-ons as share of subscription MRR, year 3", 0.14, None, "Adoption ~60% by year 3"),
    ("HOSTING", None, None, None),
    ("Azure hosting, £ per firm per month", 40, None, "Franck: ~£2,000/month for ~50 organisations; multi-tenant, scales roughly linearly"),
    ("Azure hosting floor, £ per month", 500, None, "Fixed cost of the environment at low customer counts"),
    ("ROUND", None, None, None),
    ("Round received in month", 1, None, "Model assumes close at month 1; shift to test a later close"),
    ("Option A round, £", 500000, None, "Deck slide 16"),
    ("Option B round, £", 1000000, None, "Deck slide 16"),
    ("TEAM AND SPEND — Option A (£ per year unless stated)", None, None, None),
    ("Two developers, cost per year (both)", 160000, None, "£200k over 15 months on the deck"),
    ("Developers start month", 3, None, "Within 60 days of close"),
    ("Sales hire start month", 2, None, ""),
    ("Integrations budget (Sage API, telematics, stores), £ total", 50000, None, "Spread evenly over months 3–12"),
    ("Marketing, cost per year — Option A", 40000, None, "£50k over 15 months: CampaignBuilder outbound, trade bodies, events"),
    ("Marketing, cost per year — Option B", 360000, None, "£450k over 15 months: adds national paid search, content, events"),
    ("Marketing start month", 3, None, ""),
    ("Onboarding / customer-success hire, cost per year — Option B only", 40000, None, "£50k over 15 months; starts month 4"),
    ("Onboarding hire start month (B)", 4, None, ""),
    ("Legal, finance, contingency, £ per month — Option A", 2500, None, "£60k over 15 months less hosting (modelled separately)"),
    ("Legal, finance, contingency, £ per month — Option B", 5500, None, "£110k over 15 months less hosting"),
    ("Founder salary, £ per year", 48000, None, "Franck: £4,000 a month from close; £60k over the 15-month budget on deck slide 16"),
    ("SCALING WITH VOLUME (drives years 2–3; the deck says a seed round in year 2 funds them)", None, None, None),
    ("New firms closed per salesperson per month, at full ramp", 8, None, "Franck's own target for the first hire is 8–9 a month; extra salespeople are added when the ramp needs them"),
    ("Salesperson cost per year, each", 64000, None, "Same basis as the first hire"),
    ("Set-ups delivered per onboarding person per month", 6, None, "Each set-up is about a week of work (RAMS templates, Sage mapping, data import, two training sessions)"),
    ("Set-ups the founder delivers himself per month", 3, None, "Set to 0 once the founder is out of onboarding"),
    ("Onboarding person cost per year, each", 40000, None, ""),
    ("Support staff: firms served per person (from month 13)", 100, None, "Ongoing support and account management once the base is large"),
    ("Support staff cost per year, each", 40000, None, ""),
    ("Additional developer from month 13, cost per year", 80000, None, "Third developer as the base grows"),
    ("Marketing spend per new firm from month 13, £", 1500, None, "Paid acquisition cost per firm once referrals and the first campaigns are exhausted; replaces the fixed year-1 budget"),
]
r = 4
ref = {}
for label, val, _, note in rows:
    A.cell(r, 1, label)
    if val is None:
        A.cell(r, 1).font = H2; A.cell(r, 1).fill = NAVY
        for c in range(2, 5): A.cell(r, c).fill = NAVY
    else:
        c = A.cell(r, 2, val); c.font = BLUE; c.fill = YEL
        if isinstance(val, float) and val <= 1: c.number_format = PCT
        elif isinstance(val, date): c.number_format = "mmm yyyy"
        elif isinstance(val, int) and val >= 1000: c.number_format = GBP
        A.cell(r, 4, note).font = Font(name=F, italic=True, color="555555")
        ref[label] = f"Assumptions!$B${r}"
    r += 1

# New-firms schedule
r += 1
A.cell(r, 1, "NEW FIRMS SIGNED PER MONTH (inputs; the ramp on deck slides 14–15)").font = H2
for c in range(1, 40): A.cell(r, c).fill = NAVY
sched_hdr = r + 1
A.cell(sched_hdr, 1, "Month").font = BOLD
for m in range(1, 37):
    A.cell(sched_hdr, 2 + m, m).font = BOLD
    A.column_dimensions[L(2 + m)].width = 6
A.column_dimensions["C"].width = 14
schedA = [0,0,0,2,3,4,5,5,6,7,8,9]+[9,10,10,11,11,12,12,13,13,14,14,15]+[15,16,16,17,17,18,18,19,19,20,20,21]
schedB = [0,0,0,3,5,7,8,9,11,13,14,16]+[17,19,20,22,23,24,25,26,27,28,29,30]+[31,32,33,34,35,36,37,38,39,40,41,42]
rowA, rowB = sched_hdr + 1, sched_hdr + 2
A.cell(rowA, 1, "Option A — referrals + one trades salesperson").font = BOLD
A.cell(rowB, 1, "Option B — plus national marketing + onboarding hire").font = BOLD
for m in range(36):
    for rr, s in ((rowA, schedA), (rowB, schedB)):
        c = A.cell(rr, 3 + m, s[m]); c.font = BLUE; c.fill = YEL
A.cell(rowB + 1, 1, "Salesperson starts month 2, first closes month 4 on a 30–90 day cycle; 15%/month growth from month 7 tapering in year 2. Year 2–3 pace requires the second and third salespeople above.").font = Font(name=F, italic=True, color="555555")
A.cell(rowB + 3, 1, "Legend: blue on yellow = input you can change; black = formula; green = value pulled from another sheet.").font = Font(name=F, italic=True)
A.freeze_panes = "B4"

# ---------------- Option sheets ----------------
def option_sheet(name, opt):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 44
    ws["A1"] = f"{name} — monthly model, 36 months"; ws["A1"].font = H1
    ws["A2"] = "All cells are formulas driven by the Assumptions sheet. £ excluding VAT."; ws["A2"].font = Font(name=F, italic=True)
    hdr = 4
    ws.cell(hdr, 1, "Month").font = BOLD
    ws.cell(hdr + 1, 1, "Calendar month").font = BOLD
    labels = [
        ("CUSTOMERS", None),
        ("Firms at start of month", "start"),
        ("New firms signed", "new"),
        ("Firms churned", "churn"),
        ("Firms at end of month", "end"),
        ("RECURRING REVENUE", None),
        ("Subscription MRR, £", "base"),
        ("Add-on MRR, £", "addon"),
        ("Total MRR, £", "mrr"),
        ("ARR run-rate (MRR × 12), £", "arr"),
        ("ONE-OFF", None),
        ("Set-up fees, £", "setup"),
        ("Total revenue in month, £", "rev"),
        ("COSTS", None),
        ("Developers", "dev"),
        ("Sales (headcount scales with new firms)", "sales"),
        ("Integrations", "integ"),
        ("Marketing", "mkt"),
        ("Onboarding and support (scales with volume)", "onb"),
        ("Azure hosting", "host"),
        ("Legal, finance, contingency", "legal"),
        ("Founder salary", "founder"),
        ("Total costs, £", "cost"),
        ("CASH", None),
        ("Net cash flow in month, £", "net"),
        ("Round received, £", "round"),
        ("Cash balance at end of month, £", "cash"),
    ]
    R = {}
    rr = hdr + 2
    for lab, key in labels:
        ws.cell(rr, 1, lab)
        if key is None:
            ws.cell(rr, 1).font = H2
            for c in range(1, 39): ws.cell(rr, c).fill = NAVY
        else:
            R[key] = rr
            ws.cell(rr, 1).font = BOLD if key in ("end", "mrr", "rev", "cost", "cash", "arr") else BLACK
        rr += 1
    sched_row = rowA if opt == "A" else rowB
    rnd = ref["Option A round, £"] if opt == "A" else ref["Option B round, £"]
    mkt = ref["Marketing, cost per year — Option A"] if opt == "A" else ref["Marketing, cost per year — Option B"]
    legal = ref["Legal, finance, contingency, £ per month — Option A"] if opt == "A" else ref["Legal, finance, contingency, £ per month — Option B"]
    for m in range(1, 37):
        col = 1 + m; c = L(col); p = L(col - 1)
        ws.column_dimensions[c].width = 12
        ws.cell(hdr, col, m).font = BOLD
        ws.cell(hdr + 1, col, f"=EDATE({ref['Model start (month 1)']},{c}{hdr}-1)").number_format = "mmm yy"
        ws.cell(hdr + 1, col).font = GREEN
        # customers
        ws.cell(R["start"], col, f"={ref['Existing firms at month 0']}" if m == 1 else f"={p}{R['end']}")
        ws.cell(R["new"], col, f"=ROUND(Assumptions!{L(2 + m)}{sched_row}*{ref['Sales achievement vs plan (sensitivity)']},0)").font = GREEN
        ws.cell(R["churn"], col, f"={c}{R['start']}*{ref['Monthly logo churn']}")
        ws.cell(R["end"], col, f"={c}{R['start']}-{c}{R['churn']}+{c}{R['new']}")
        # base MRR
        prev_base = f"{ref['Existing MRR at month 0 (CBES net £440 + Gaswise £2,600), £']}" if m == 1 else f"{p}{R['base']}"
        ws.cell(R["base"], col,
                f"=({prev_base}*(1-{ref['Monthly logo churn']})+{c}{R['new']}*{ref['Average new contract, £ per month at list']}"
                f"+IF({c}{hdr}={ref['CBES discount ends: month']},{ref['CBES MRR step-up when discount ends, £']},0))"
                f"*IF(OR({c}{hdr}=13,{c}{hdr}=25),1+{ref['Price uplift at each annual renewal']},1)")
        share = (f"IF({c}{hdr}<=12,{ref['Add-ons as share of subscription MRR, year 1']},"
                 f"IF({c}{hdr}<=24,{ref['Add-ons as share of subscription MRR, year 2']},{ref['Add-ons as share of subscription MRR, year 3']}))")
        ws.cell(R["addon"], col, f"={c}{R['base']}*{share}")
        ws.cell(R["mrr"], col, f"={c}{R['base']}+{c}{R['addon']}")
        ws.cell(R["arr"], col, f"={c}{R['mrr']}*12")
        ws.cell(R["setup"], col, f"={c}{R['new']}*{ref['Set-up fee per new firm, £ one-off']}")
        ws.cell(R["rev"], col, f"={c}{R['mrr']}+{c}{R['setup']}")
        # costs
        ws.cell(R["dev"], col, f"=IF({c}{hdr}>={ref['Developers start month']},{ref['Two developers, cost per year (both)']}/12,0)"
                f"+IF({c}{hdr}>=13,{ref['Additional developer from month 13, cost per year']}/12,0)")
        ws.cell(R["sales"], col, f"=IF({c}{hdr}>={ref['Sales hire start month']},MAX(1,ROUNDUP({c}{R['new']}/{ref['New firms closed per salesperson per month, at full ramp']},0)),0)"
                f"*{ref['Salesperson cost per year, each']}/12")
        ws.cell(R["integ"], col, f"=IF(AND({c}{hdr}>=3,{c}{hdr}<=12),{ref['Integrations budget (Sage API, telematics, stores), £ total']}/10,0)")
        ws.cell(R["mkt"], col, f"=IF({c}{hdr}<=12,IF({c}{hdr}>={ref['Marketing start month']},{mkt}/12,0),{c}{R['new']}*{ref['Marketing spend per new firm from month 13, £']})")
        onb_first = (f"IF({c}{hdr}>={ref['Onboarding hire start month (B)']},{ref['Onboarding / customer-success hire, cost per year — Option B only']}/12,0)"
                     if opt == "B" else "0")
        need = f"ROUNDUP(MAX(0,{c}{R['new']}-{ref['Set-ups the founder delivers himself per month']})/{ref['Set-ups delivered per onboarding person per month']},0)"
        extra = f"MAX(0,{need}-{'1' if opt == 'B' else '0'})" if True else need
        ws.cell(R["onb"], col, f"={onb_first}+IF({c}{hdr}>={ref['Onboarding hire start month (B)']},{extra},0)*{ref['Onboarding person cost per year, each']}/12"
                f"+IF({c}{hdr}>=13,ROUNDUP({c}{R['start']}/{ref['Support staff: firms served per person (from month 13)']},0)*{ref['Support staff cost per year, each']}/12,0)")
        ws.cell(R["host"], col, f"=MAX({ref['Azure hosting floor, £ per month']},{c}{R['end']}*{ref['Azure hosting, £ per firm per month']})")
        ws.cell(R["legal"], col, f"={legal}")
        ws.cell(R["founder"], col, f"={ref['Founder salary, £ per year']}/12")
        ws.cell(R["cost"], col, f"=SUM({c}{R['dev']}:{c}{R['founder']})")
        ws.cell(R["net"], col, f"={c}{R['rev']}-{c}{R['cost']}")
        ws.cell(R["round"], col, f"=IF({c}{hdr}={ref['Round received in month']},{rnd},0)")
        ws.cell(R["cash"], col, f"={c}{R['net']}+{c}{R['round']}" if m == 1 else f"={p}{R['cash']}+{c}{R['net']}+{c}{R['round']}")
        for k in ("start", "new", "churn", "end"):
            ws.cell(R[k], col).number_format = NUM1 if k in ("churn", "start", "end") else NUM
        for k in ("base", "addon", "mrr", "arr", "setup", "rev", "dev", "sales", "integ", "mkt", "onb", "host", "legal", "founder", "cost", "net", "round", "cash"):
            ws.cell(R[k], col).number_format = GBP
    for k in ("end", "mrr", "rev", "cost", "cash"):
        for col in range(2, 38): ws.cell(R[k], col).font = BOLD
    ws.freeze_panes = "B6"
    return ws, R, hdr

wsA, RA, hdrA = option_sheet("Option A", "A")
wsB, RB, hdrB = option_sheet("Option B", "B")

# ---------------- Summary ----------------
S = wb.create_sheet("Summary", 0)
S.column_dimensions["A"].width = 44
for c in "BCDEFGH": S.column_dimensions[c].width = 15
S["A1"] = "Jobs Tracker — summary against the investor deck"; S["A1"].font = H1
S["A2"] = "Model figures are formulas from the option sheets; 'Deck' columns are the rounded figures printed on slides 14–16 (hardcoded for comparison)."; S["A2"].font = Font(name=F, italic=True)

def block(top, title, opt, ws, R, hdr, deck):
    S.cell(top, 1, title).font = H2
    for c in range(1, 9): S.cell(top, c).fill = NAVY
    heads = ["", "M4", "M6", "M9", "M12", "Year 1", "Year 2", "Year 3"]
    for i, h in enumerate(heads):
        S.cell(top + 1, 1 + i, h).font = BOLD; S.cell(top + 1, 1 + i).fill = GREY
    def colm(m): return L(1 + m)
    rows = [
        ("New firms signed in month / average per month in year", lambda m: f"='{ws.title}'!{colm(m)}{R['new']}", lambda y: f"=AVERAGE('{ws.title}'!{colm(y*12-11)}{R['new']}:{colm(y*12)}{R['new']})", NUM1),
        ("Firms on the platform at end", lambda m: f"='{ws.title}'!{colm(m)}{R['end']}", lambda y: f"='{ws.title}'!{colm(y*12)}{R['end']}", NUM),
        ("Total MRR at end, £", lambda m: f"='{ws.title}'!{colm(m)}{R['mrr']}", lambda y: f"='{ws.title}'!{colm(y*12)}{R['mrr']}", GBP),
        ("ARR run-rate at end, £", lambda m: f"='{ws.title}'!{colm(m)}{R['arr']}", lambda y: f"='{ws.title}'!{colm(y*12)}{R['arr']}", GBP),
        ("Set-up fees in year, £", None, lambda y: f"=SUM('{ws.title}'!{colm(y*12-11)}{R['setup']}:{colm(y*12)}{R['setup']})", GBP),
        ("Revenue in year incl. set-up, £", None, lambda y: f"=SUM('{ws.title}'!{colm(y*12-11)}{R['rev']}:{colm(y*12)}{R['rev']})", GBP),
        ("Costs in year, £", None, lambda y: f"=SUM('{ws.title}'!{colm(y*12-11)}{R['cost']}:{colm(y*12)}{R['cost']})", GBP),
        ("Cash at end, £", lambda m: f"='{ws.title}'!{colm(m)}{R['cash']}", lambda y: f"='{ws.title}'!{colm(y*12)}{R['cash']}", GBP),
    ]
    r = top + 2
    for lab, fm, fy, fmt in rows:
        S.cell(r, 1, lab)
        for i, m in enumerate((4, 6, 9, 12)):
            if fm: c = S.cell(r, 2 + i, fm(m)); c.number_format = fmt; c.font = GREEN
        for y in (1, 2, 3):
            c = S.cell(r, 5 + y, fy(y)); c.number_format = fmt; c.font = GREEN
        r += 1
    S.cell(r, 1, "Lowest cash balance in 36 months, £ / month").font = BOLD
    last = L(37)
    S.cell(r, 2, f"=MIN('{ws.title}'!B{R['cash']}:{last}{R['cash']})").number_format = GBP
    S.cell(r, 3, f"=MATCH(B{r},'{ws.title}'!B{R['cash']}:{last}{R['cash']},0)").number_format = NUM
    S.cell(r, 4, "← month index; negative means a seed round is needed before that month").font = Font(name=F, italic=True, color="555555")
    r += 1
    S.cell(r, 1, "Months of runway on this round alone (first month cash < 0)").font = BOLD
    S.cell(r, 2, f'=IFERROR(MATCH(TRUE,INDEX(\'{ws.title}\'!B{R["cash"]}:{last}{R["cash"]}<0,0),0)-1,">36")').number_format = NUM
    r += 2
    S.cell(r, 1, "Deck (slides 14–16) shows:").font = BOLD
    for i, h in enumerate(["", "M4", "M6", "M9", "M12", "Year 1", "Year 2", "Year 3"]):
        S.cell(r, 1 + i).fill = GREY;
        if i: S.cell(r, 1 + i, h).font = BOLD
    r += 1
    for lab, vals, fmt in deck:
        S.cell(r, 1, lab)
        for i, v in enumerate(vals):
            if v is not None:
                c = S.cell(r, 2 + i, v); c.font = BLUE; c.number_format = fmt
        r += 1
    return r + 1

deckA = [("Firms on the platform", [4, 11, 27, 51, 50, 180, 364], NUM),
         ("MRR, £", [5000, 13000, 32000, 60000, None, None, None], GBP),
         ("ARR run-rate, £", [None, None, None, None, 700000, 2800000, 6000000], GBP),
         ("Set-up fees in year, £", [None, None, None, None, 195000, None, None], GBP),
         ("Revenue in year, £", [None, None, None, None, 500000, 2500000, 5500000], GBP)]
deckB = [("Firms on the platform", [5, 17, 45, 88, 85, 348, 718], NUM),
         ("MRR, £", [7000, 20000, 53000, 100000, None, None, None], GBP),
         ("ARR run-rate, £", [None, None, None, None, 1200000, 5400000, 11800000], GBP),
         ("Set-up fees in year, £", [None, None, None, None, 345000, None, None], GBP),
         ("Revenue in year, £", [None, None, None, None, 800000, 4700000, 10700000], GBP)]
nxt = block(4, "OPTION A — £500k: referrals + one trades salesperson", "A", wsA, RA, hdrA, deckA)
nxt = block(nxt, "OPTION B — £1m: plus national marketing and an onboarding hire", "B", wsB, RB, hdrB, deckB)
S.cell(nxt, 1, "Targets, not forecasts: there is no repeatable sales channel yet. Change the new-firms rows on the Assumptions sheet to test other ramps.").font = Font(name=F, italic=True)
S.freeze_panes = "B4"

for ws in wb.worksheets: style_all(ws)
out = "/tmp/claude-0/-home-claude/e6a6822b-765a-572c-803c-ac05cb6d9117/scratchpad/model/JobsTracker_Financial_Model_2026-10-18.xlsx"
wb.save(out); print(out)
